import concurrent.futures
import shutil
import tempfile
import unittest
from pathlib import Path
import openpyxl
import server

class RegistrationTests(unittest.TestCase):
    def setUp(self):
        server.initialize_excel()
        self.tmp=tempfile.TemporaryDirectory(); self.book=Path(self.tmp.name)/'test.xlsx'; shutil.copy2(server.BOOK,self.book)
        wb=openpyxl.load_workbook(self.book)
        for s in wb:
            if s.title!='Allocations':
                for r in s.iter_rows(min_row=2,max_col=12):
                    for cell in r:cell.value=None
        wb.save(self.book);wb.close()
    def tearDown(self): self.tmp.cleanup()
    def payload(self,i=1):
        return dict(name='Test Student',register=f'TEST000{i}',email=f'test{i}@mkce.ac.in',phone='9876543210',department='Civil Engineering',year='1',japanese='No',club='Google Club')
    def test_save_and_duplicate(self):
        result=server.register(self.payload(),self.book)
        wb=openpyxl.load_workbook(self.book); self.assertEqual(wb['Google Club']['K2'].value,result['id']); self.assertEqual(wb['Google Club']['B2'].value,'TEST0001'); self.assertEqual(wb['Google Club']['F2'].value,1); self.assertEqual(wb['Google Club']['J2'].value,'Registered'); self.assertEqual(wb['Google Club']['A1'].value,'Name'); self.assertIsNotNone(wb['Google Club']['L2'].value); self.assertEqual(server.read_state(wb)['allocations'][8]['registered'],1); wb.close()
        p=self.payload(); p['club']='Mobile App Club'
        with self.assertRaisesRegex(ValueError,'already registered'): server.register(p,self.book)
        p=self.payload(2); p['email']='test1@mkce.ac.in'
        with self.assertRaisesRegex(ValueError,'already registered'): server.register(p,self.book)
    def test_last_seat_concurrent(self):
        def attempt(i):
            try: server.register(self.payload(i),self.book); return True
            except ValueError: return False
        with concurrent.futures.ThreadPoolExecutor(2) as pool: self.assertEqual(sum(pool.map(attempt,[1,2])),1)
    def test_invalid_and_formula_text(self):
        p=self.payload(); p['email']='test@gmail.com'
        with self.assertRaisesRegex(ValueError,'college email'): server.register(p,self.book)
        p=self.payload(); p['name']='=HYPERLINK("evil")'
        p['aws']='Yes'
        server.register(p,self.book); wb=openpyxl.load_workbook(self.book); self.assertEqual(wb['Google Club']['A2'].data_type,'s'); self.assertIsNone(wb['Google Club']['H2'].value); wb.close()
    def test_japanese_and_fourth_year_rules(self):
        p=self.payload();p['year']='4'
        with self.assertRaisesRegex(ValueError,'1st, 2nd and 3rd'):server.register(p,self.book)
        p=self.payload();p['japanese']='Yes'
        with self.assertRaisesRegex(ValueError,'Japanese Club'):server.register(p,self.book)
        p['club']='Japanese Club';server.register(p,self.book)
        p=self.payload(2);p.update(japanese='Yes',club='Japanese Club');server.register(p,self.book)
        wb=openpyxl.load_workbook(self.book);self.assertEqual(wb['Japanese Club']['B3'].value,'TEST0002');wb.close()
        p=self.payload(3);p['club']='Japanese Club'
        with self.assertRaisesRegex(ValueError,'regular club'):server.register(p,self.book)

if __name__=='__main__': unittest.main()
